"""Server-side file inspection, so the browser never parses a PDF or workbook again."""

from __future__ import annotations

import io
import logging
import mimetypes
import os

_log = logging.getLogger("p0.file_metadata")

PDF_EXTENSIONS = {".pdf"}
SPREADSHEET_EXTENSIONS = {".xlsx", ".xls", ".xlsb", ".csv", ".tsv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}

_MAX_SHEET_SCAN = 25


def content_type_for(filename: str) -> str:
    """Best-effort MIME type from the filename."""
    guessed, _ = mimetypes.guess_type(filename or "")
    return guessed or "application/octet-stream"


def _extension(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _pdf_metadata(payload: bytes) -> dict:
    """Page count and whether a text layer exists — drives the OCR-vs-extract decision."""
    try:
        import fitz
    except ImportError:
        _log.debug("[file_metadata] PyMuPDF not installed — PDF metadata skipped")
        return {}
    try:
        with fitz.open(stream=payload, filetype="pdf") as doc:
            pages = []
            has_text = False
            for index in range(doc.page_count):
                page = doc[index]
                text = page.get_text().strip()
                if text:
                    has_text = True
                rect = page.rect
                pages.append(
                    {
                        "page_number": index + 1,
                        "width_px": int(rect.width),
                        "height_px": int(rect.height),
                        "rotation": int(page.rotation or 0),
                        "has_text": bool(text),
                    }
                )
            return {
                "kind": "pdf",
                "page_count": doc.page_count,
                "has_text_layer": has_text,
                "pages": pages,
            }
    except Exception as exc:
        _log.warning("[file_metadata] PDF inspection failed: %s", exc)
        return {}


def _spreadsheet_metadata(payload: bytes, extension: str) -> dict:
    """Sheet names, header row and row counts, without loading the whole workbook."""
    if extension in (".csv", ".tsv"):
        return _delimited_metadata(payload, extension)
    try:
        import openpyxl
    except ImportError:
        _log.debug("[file_metadata] openpyxl not installed — workbook metadata skipped")
        return {}
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(payload), read_only=True, data_only=True
        )
        sheets = []
        for worksheet in list(workbook.worksheets)[:_MAX_SHEET_SCAN]:
            header = []
            try:
                first = next(worksheet.iter_rows(values_only=True), None)
                header = [str(c).strip() for c in (first or []) if c is not None]
            except Exception:
                header = []
            sheets.append(
                {
                    "name": worksheet.title,
                    "row_count": int(worksheet.max_row or 0),
                    "column_count": int(worksheet.max_column or 0),
                    "header_row": 1 if header else None,
                    "columns": header,
                }
            )
        workbook.close()
        return {"kind": "spreadsheet", "sheets": sheets}
    except Exception as exc:
        _log.warning("[file_metadata] workbook inspection failed: %s", exc)
        return {}


def _delimited_metadata(payload: bytes, extension: str) -> dict:
    """Header and row count for csv/tsv, read as text."""
    delimiter = "\t" if extension == ".tsv" else ","
    try:
        text = payload.decode("utf-8", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        header = [c.strip() for c in lines[0].split(delimiter)] if lines else []
        return {
            "kind": "spreadsheet",
            "sheets": [
                {
                    "name": "(csv)",
                    "row_count": max(0, len(lines) - 1),
                    "column_count": len(header),
                    "header_row": 1 if header else None,
                    "columns": header,
                }
            ],
        }
    except Exception as exc:
        _log.warning("[file_metadata] delimited inspection failed: %s", exc)
        return {}


def _image_metadata(payload: bytes) -> dict:
    """Pixel dimensions and DPI where the image carries them."""
    try:
        from PIL import Image
    except ImportError:
        return {}
    try:
        with Image.open(io.BytesIO(payload)) as image:
            dpi = image.info.get("dpi")
            return {
                "kind": "image",
                "width_px": image.width,
                "height_px": image.height,
                "dpi": int(dpi[0]) if dpi else None,
            }
    except Exception as exc:
        _log.warning("[file_metadata] image inspection failed: %s", exc)
        return {}


def inspect(filename: str, payload: bytes) -> dict:
    """Everything the server can learn about an uploaded file, without the browser."""
    extension = _extension(filename)
    metadata = {
        "kind": "unknown",
        "extension": extension.lstrip("."),
        "content_type": content_type_for(filename),
        "size_bytes": len(payload or b""),
    }
    if not payload:
        return metadata

    if extension in PDF_EXTENSIONS:
        metadata.update(_pdf_metadata(payload))
    elif extension in SPREADSHEET_EXTENSIONS:
        metadata.update(_spreadsheet_metadata(payload, extension))
    elif extension in IMAGE_EXTENSIONS:
        metadata.update(_image_metadata(payload))
    return metadata


def flow_fields(metadata: dict) -> dict:
    """The subset of inspect() that flow_state stores as columns."""
    if not metadata:
        return {}
    fields = {"content_type": metadata.get("content_type")}
    if metadata.get("kind") == "pdf":
        fields["page_count"] = metadata.get("page_count")
        fields["has_text_layer"] = metadata.get("has_text_layer")
    elif metadata.get("kind") == "spreadsheet":
        sheets = metadata.get("sheets") or []
        fields["sheet_names"] = [s.get("name") for s in sheets]
        total = sum(int(s.get("row_count") or 0) for s in sheets)
        if total:
            fields["rows_in"] = total
    return {k: v for k, v in fields.items() if v is not None}
